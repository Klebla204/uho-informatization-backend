from dataclasses import dataclass, field
from datetime import datetime
import uuid

from django.db import transaction
from openpyxl import load_workbook

from .models import Autor, Biblioteca, Categoria, Editorial, Ejemplar, Libro

MAX_IMPORT_ROWS = 1000
REQUIRED_COLUMNS = {
    "titulo",
    "autor_principal",
    "editorial",
    "año_edicion",
    "categoria",
    "biblioteca",
    "cantidad_ejemplares",
    "activo",
}


@dataclass
class ImportResult:
    processed: int = 0
    created: int = 0
    updated: int = 0
    copies_created: int = 0
    errors: list[dict] = field(default_factory=list)


def _text(value):
    return "" if value is None else str(value).strip()


def _parse_active(value):
    value = _text(value).upper()
    if value not in {"SI", "NO"}:
        raise ValueError("activo debe ser SI o NO")
    return value == "SI"


def _parse_year(value):
    try:
        year = int(value)
    except (TypeError, ValueError):
        raise ValueError("año_edicion debe ser un número entero")
    if year < 1800 or year > datetime.now().year:
        raise ValueError("año_edicion está fuera del rango permitido")
    return year


def _parse_positive_integer(value, field_name):
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError(f"{field_name} debe ser un número entero")
    if number < 1:
        raise ValueError(f"{field_name} debe ser mayor o igual que 1")
    return number


def _find_or_create_author(value):
    parts = _text(value).split(",", 1)
    if len(parts) != 2 or not all(parts):
        raise ValueError("autor_principal debe tener formato Apellidos, Nombre")
    apellidos, nombre = (part.strip() for part in parts)
    return Autor.objects.get_or_create(nombre=nombre, apellidos=apellidos)[0]


def _find_or_create_category(value, parent_name=""):
    category, _ = Categoria.objects.get_or_create(nombre=_text(value), defaults={"activa": True})
    if parent_name:
        parent, _ = Categoria.objects.get_or_create(nombre=_text(parent_name), defaults={"activa": True})
        if category.categoria_padre_id != parent.id:
            category.categoria_padre = parent
            category.save(update_fields=["categoria_padre"])
    return category


def _validate_headers(headers):
    missing = REQUIRED_COLUMNS.difference(set(headers))
    if missing:
        raise ValueError(f"faltan columnas obligatorias: {', '.join(sorted(missing))}")


def import_collection(file_object, user=None):
    workbook = load_workbook(file_object, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = [_text(value) for value in next(rows, ())]
    _validate_headers(headers)

    result = ImportResult()
    for row_number, values in enumerate(rows, start=2):
        if result.processed >= MAX_IMPORT_ROWS:
            result.errors.append({"fila": row_number, "motivo": "el límite es de 1000 filas"})
            continue
        if not any(value is not None for value in values):
            continue

        result.processed += 1
        data = dict(zip(headers, values))
        try:
            with transaction.atomic():
                title = _text(data.get("titulo"))
                if not title or len(title) > 500:
                    raise ValueError("titulo es obligatorio y no puede superar 500 caracteres")
                isbn = _text(data.get("isbn")) or None
                editorial, _ = Editorial.objects.get_or_create(nombre=_text(data["editorial"]), defaults={"activa": True})
                category = _find_or_create_category(data["categoria"], data.get("subcategoria"))
                author = _find_or_create_author(data["autor_principal"])
                copies = _parse_positive_integer(data["cantidad_ejemplares"], "cantidad_ejemplares")
                library = Biblioteca.objects.get(nombre=_text(data["biblioteca"]), activa=True)
                active = _parse_active(data["activo"])
                defaults = {
                    "titulo": title,
                    "anio_edicion": _parse_year(data["año_edicion"]),
                    "editorial": editorial,
                    "edicion": _text(data.get("edicion")),
                    "resumen": _text(data.get("resumen")),
                    "portada_url": _text(data.get("portada_url")),
                    "activo": active,
                }
                if isbn:
                    book, created = Libro.objects.update_or_create(isbn=isbn, defaults=defaults)
                else:
                    book = Libro.objects.create(**defaults, isbn=None)
                    created = True
                if created:
                    result.created += 1
                else:
                    result.updated += 1
                book.autores.add(author)
                book.categorias.add(category)
                for index in range(copies):
                    Ejemplar.objects.create(
                        libro=book,
                        biblioteca=library,
                        codigo_ejemplar=f"{book.id}-{library.id}-{uuid.uuid4().hex[:12]}",
                        ubicacion_fisica=_text(data.get("ubicacion_fisica")),
                    )
                    result.copies_created += 1
        except Exception as error:
            result.errors.append({"fila": row_number, "motivo": str(error)})

    return result
